# -*- coding: utf-8 -*-
"""★ 1.9.3 第 40 轮：**测试自造 fixture**，不再依赖 `data/桌宠私有data/` 里的用户文件。

为什么（真机教训 2026-10-04）：
  用户让桌宠"整理私有文件夹"，模型用 `shutil.rmtree` 把
  `data/桌宠私有data/` 清空（301 文件 / 102.5 MB）——
  **全量回归立刻出现 4 处红**，全是"套件依赖用户文件"：
    · `_test_bg_thread_ui_182` / `_test_attach_persist_184` / `_test_file_card_182`
      依赖 `office/pptx/deepseek_intro.pptx`
    · `_test_pptx_render_193` 依赖 `迷迭香介绍.pptx`
    · `_test_a190` 的 whisper 转写断言依赖音频文件
    · `_test_plugins` 断言用户的 `office` 目录里"至少 3 个文件"
  `data/桌宠私有data/` 是**用户的地盘**，随时会被清 —— 测试必须自给自足。

用法：
    from _fixtures_193 import make_pptx, make_wav, fixture_dir
    p = make_pptx(fixture_dir("xxx") / "sample.pptx")
"""
import struct
import wave
from pathlib import Path

BASE = Path(__file__).resolve().parents[2]

# 所有 fixture 都放这里（在 tools/dev_tests 下，随 `_animsrc_193` 一起被 gitignore）
FIXTURE_ROOT = Path(__file__).resolve().parent / "_fixtures_193"


def fixture_dir(name: str) -> Path:
    """给某个套件一个自己的 fixture 目录（存在即返回）。"""
    d = FIXTURE_ROOT / name
    d.mkdir(parents=True, exist_ok=True)
    return d


def make_pptx(dest, pages: int = 2, title: str = "测试演示文稿") -> Path:
    """造一个**真** .pptx（python-pptx 生成，内容无所谓，够真就行）。"""
    dest = Path(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)
    if dest.exists() and dest.stat().st_size > 0:
        return dest
    from pptx import Presentation
    from pptx.util import Inches, Pt
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    s = prs.slides.add_slide(blank)
    tb = s.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11), Inches(1.4))
    tb.text_frame.text = title
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(40)
    for i in range(1, max(1, pages)):
        sp = prs.slides.add_slide(blank)
        t2 = sp.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11), Inches(1.2))
        t2.text_frame.text = f"第 {i + 1} 页"
    prs.save(str(dest))
    return dest


def make_pdf(dest, pages: int = 8, sparse_page: int = 2) -> Path:
    """造一个**真** PDF（PyMuPDF/fitz 生成）—— 供 PDF 端到端断言用。

    ★ 1.9.4 阶段2 为什么补这个：
      `_test_a190` 里有一块 "PDF 端到端两项" 依赖
      `data/downloads/jmcomic/PDFs/*.pdf`（**用户下载来的**）。
      用户在模拟全新状态时删掉了 `data/`，那块就**静默跳过**（打印 NOTE，但计数从
      463 掉到 461）—— 也就是说**新用户那里这两项永远不跑**。
      这与"独立运行 / 新用户流程"直接冲突，所以改成自造。

    形状要求（照 `file_view._view_pdf` 的判据反推）：
      · **总页数 > 5** —— 断言里核对 `pages_total > 5`；
      · **`sparse_page` 那一页的文字 < 30 字** —— `_view_pdf` 的逻辑是
        "该页文本 < 30 字就渲染成图走**视觉**"，所以那页**故意留白**。
    """
    dest = Path(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)
    if dest.exists() and dest.stat().st_size > 0:
        return dest
    import fitz                                    # PyMuPDF；项目读 PDF 用的就是它
    doc = fitz.open()
    for i in range(1, max(6, pages) + 1):
        page = doc.new_page()
        if i == sparse_page:
            continue                               # ★ 留白 → 触发"文本太少走视觉"
        page.insert_text((72, 72),
                         f"Fixture 第 {i} 页。" + "这是用于回归的测试文本。" * 6)
    doc.save(str(dest))
    doc.close()
    return dest


def make_wav(dest, seconds: float = 1.0, freq: int = 440,
             rate: int = 16000) -> Path:
    """造一个**真**的 16kHz 单声道 PCM wav（给转写/音频链路当输入）。"""
    import math
    dest = Path(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)
    if dest.exists() and dest.stat().st_size > 0:
        return dest
    n = int(rate * seconds)
    with wave.open(str(dest), "wb") as w:
        w.setnchannels(1)
        w.setsampwidth(2)
        w.setframerate(rate)
        frames = bytearray()
        for i in range(n):
            v = int(12000 * math.sin(2 * math.pi * freq * i / rate))
            frames += struct.pack("<h", v)
        w.writeframes(bytes(frames))
    return dest


def make_speech_wav(dest, text: str = "这是一段用于回归测试的语音。") -> Path | None:
    """用 Windows 自带的 SAPI（System.Speech）合成一段**真语音** wav。

    ★ 为什么需要：whisper 真实转写那条断言要的是"有内容的语音"，
      正弦波转不出文字（会变成假失败）。用 SAPI 合成 → 自带、零新依赖。
    ★ 拿不到（非 Windows / 没装语音 / 合成失败）就返回 None，
      调用方应当**明确 SKIP 并打印原因**，而不是判失败。
    """
    dest = Path(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)
    if dest.exists() and dest.stat().st_size > 1000:
        return dest
    try:
        import subprocess
        ps = (
            "Add-Type -AssemblyName System.Speech;"
            "$s = New-Object System.Speech.Synthesis.SpeechSynthesizer;"
            f"$s.SetOutputToWaveFile('{dest}');"
            f"$s.Speak('{text}');"
            "$s.Dispose()"
        )
        # ★ 不要 capture_output：DSH 沙箱下"管道 stdio"会被拒（EPERM）。
        # ★★ 1.9.4：**必须加 no_window_kwargs()** —— 否则每次调用都会在屏幕上
        #   闪一个 PowerShell 黑框（用户真机反馈："屏幕不停的闪现 cmd 黑色窗口"）。
        subprocess.run(["powershell", "-NoProfile", "-Command", ps],
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                       timeout=60, **no_window_kwargs())
    except Exception:
        return None
    return dest if (dest.exists() and dest.stat().st_size > 1000) else None


def make_png(dest, size=(320, 320)) -> Path:
    """造一张真 PNG（给图片/卡片链路当输入）。"""
    dest = Path(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)
    if dest.exists() and dest.stat().st_size > 0:
        return dest
    from PySide6.QtGui import QImage, QColor
    img = QImage(size[0], size[1], QImage.Format_ARGB32)
    img.fill(QColor(40, 90, 160))
    img.save(str(dest), "PNG")
    return dest


if __name__ == "__main__":
    d = fixture_dir("smoke")
    print("pptx:", make_pptx(d / "a.pptx").stat().st_size, "字节")
    print("wav :", make_wav(d / "a.wav").stat().st_size, "字节")
    print("fixture 根目录:", FIXTURE_ROOT)


def no_window_kwargs() -> dict:
    """★ 1.9.4：起子进程时**别弹黑窗口**。

    用户真机反馈（2026-10-06）："你的这一步动作导致我的屏幕不停的闪现 cmd
    黑色窗口，持续约3秒。担心是bug。" —— 原因是 Windows 上 `subprocess` 起
    控制台程序**默认会弹窗**；回归套件里每起一个 python/powershell 子进程就闪一下。

    用法：
        subprocess.run(cmd, ..., **no_window_kwargs())

    非 Windows 返回空 dict（`creationflags` 是 Windows 专有参数）。
    """
    import sys as _sys
    if _sys.platform != "win32":
        return {}
    # CREATE_NO_WINDOW = 0x08000000（不给子进程建控制台窗口）
    return {"creationflags": 0x08000000}
