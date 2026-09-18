"""1.8.0 聊天文件链路自测：file_ops read 的多类型支持
覆盖: 纯文本 / xlsx抽取 / docx抽取 / pptx抽取 / 图片元信息 / 二进制嗅探 / 分段读取
依赖 office 测试生成的文档（不存在时自动生成）
"""
import sys
import io
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
BASE = Path(__file__).resolve().parent.parent.parent
sys.path.insert(0, str(BASE))

from tools.file_ops import file_read  # noqa: E402

PRIV = BASE / "data" / "桌宠私有data"
results = []


def check(name, ok, detail=""):
    results.append(ok)
    print(f"  [{'PASS' if ok else 'FAIL'}] {name}  {str(detail)[:80]}")


print("== 1. 纯文本读取 ==")
txt = PRIV / "_kg_file_read_test.txt"
txt.write_text("第一行\n第二行\n第三行", encoding="utf-8")
r = file_read(str(txt))
check("读取全部行", "第一行" in r and "第三行" in r, r[:60])
r2 = file_read(str(txt), offset=2, limit=1)
check("offset/limit 分段", "第二行" in r2 and "第三行" not in r2, r2[:60])

print("== 2. Excel 抽取 ==")
xlsx = PRIV / "office" / "xlsx" / "自检表格.xlsx"
if not xlsx.exists():
    try:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(["姓名", "分数"])
        ws.append(["阿米娅", 95])
        wb.save(xlsx)
    except Exception as e:
        print(f"  [SKIP] openpyxl 不可用: {e}")
else:
    # 清空重写保证断言稳定
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["姓名", "分数"])
    ws.append(["阿米娅", 95])
    wb.save(xlsx)
if xlsx.exists():
    r = file_read(str(xlsx))
    check("xlsx 抽取含表头和数值", "姓名" in r and "阿米娅" in r and "95" in r, r[:80])
    check("xlsx 标注工作表", "工作表" in r)

print("== 3. Word 抽取 ==")
docx = PRIV / "office" / "docx" / "_read_test.docx"
try:
    from docx import Document
    d = Document()
    d.add_paragraph("桌宠文件读取测试段落")
    d.save(docx)
    r = file_read(str(docx))
    check("docx 抽取段落", "桌宠文件读取测试段落" in r, r[:60])
except ImportError:
    print("  [SKIP] python-docx 不可用")

print("== 4. PPT 抽取 ==")
pptx = PRIV / "office" / "pptx" / "_read_test.pptx"
try:
    from pptx import Presentation
    p = Presentation()
    slide = p.slides.add_slide(p.slide_layouts[5])
    slide.shapes.title.text = "读取测试标题"
    p.save(pptx)
    r = file_read(str(pptx))
    check("pptx 抽取标题", "读取测试标题" in r, r[:60])
except ImportError:
    print("  [SKIP] python-pptx 不可用")

print("== 5. 图片元信息 ==")
try:
    from PIL import Image
    img = PRIV / "_read_test.png"
    Image.new("RGB", (320, 200), "red").save(img)
    r = file_read(str(img))
    check("图片返回尺寸", "320x200" in r, r[:80])
    check("图片提示走视觉通道", "多模态" in r or "视觉" in r)
    img.unlink()
except ImportError:
    print("  [SKIP] PIL 不可用")

print("== 6. 二进制嗅探 ==")
bin_f = PRIV / "_read_test.bin"
bin_f.write_bytes(b"\x00\x01\x02MZ bin\x00\x00 rest")
r = file_read(str(bin_f))
check("二进制识别为 [二进制文件]", "[二进制文件]" in r, r[:60])
check("十六进制头部", "00 01 02" in r)
bin_f.unlink()

print("== 7. 错误处理 ==")
r = file_read(str(PRIV / "_不存在_的文件.txt"))
check("不存在文件友好提示", "未找到" in r or "不存在" in r, r[:50])

txt.unlink(missing_ok=True)
docx.unlink(missing_ok=True)
pptx.unlink(missing_ok=True)

print(f"\n结果: {'ALL PASS' if all(results) else 'FAILED'} ({sum(results)}/{len(results)})")
sys.exit(0 if all(results) else 1)
