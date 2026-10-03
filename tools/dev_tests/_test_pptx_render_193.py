# -*- coding: utf-8 -*-
"""★ 1.9.3 第 21 轮：pptx 工具的 `render` 动作（调本机 PowerPoint 导出 PNG 核对排版）。

为什么要这一步（见 `PPT制作过程分析_1.9.3.md` 痛点 3）：
  以前生成完 PPT **没法"看一眼"** —— 错版、文字溢出、图片变形、字体缺失全发现不了。

本套件分两层：
  A. **纯函数层**（不依赖 Office，任何机器都能跑）：页号解析 `_parse_pages`；
  B. **真机层**（本机有 PowerPoint + pywin32 才跑）：造一份 3 页 pptx →
     `render` → 断言 PNG 真的落盘且能解码、`pages` 只导指定页、
     用户原本开着的 PowerPoint 不被 Quit。

⚠ 本套件会真的拉起一次 POWERPNT.EXE（WithWindow=False，不弹窗）；
  跑完**自己关掉**（若跑之前没开着）。
"""
import json
import os
import sys
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
ROOT = Path(__file__).resolve().parent.parent.parent
sys.path.insert(0, str(ROOT))

import tools.pptx_tool as P  # noqa: E402
from PIL import Image        # noqa: E402

PASS, FAIL = [], []


def check(name, ok, detail=""):
    (PASS if ok else FAIL).append(name)
    print(f"  {'PASS' if ok else 'FAIL'}  {name}" + (f"   {detail}" if detail else ""))


print("\n== A. 页号解析 _parse_pages（纯函数，与 Office 无关）==")
check("不填 → 全部页", P._parse_pages("", 5) == [1, 2, 3, 4, 5], str(P._parse_pages("", 5)))
check("None → 全部页", P._parse_pages(None, 3) == [1, 2, 3])
check("单页 \"3\"", P._parse_pages("3", 5) == [3], str(P._parse_pages("3", 5)))
check("区间 \"2-4\"", P._parse_pages("2-4", 5) == [2, 3, 4], str(P._parse_pages("2-4", 5)))
check("列表 \"1,4,7\"（越界 7 丢弃）", P._parse_pages("1,4,7", 5) == [1, 4])
check("中文逗号/全角横线也认",
      P._parse_pages("1，4－5", 6) == [1, 4, 5], str(P._parse_pages("1，4－5", 6)))
check("倒序区间 \"4-2\" 自动扶正", P._parse_pages("4-2", 5) == [2, 3, 4])
check("★全部越界 → 退回全部（不返回空，免得\"导了 0 张\"）",
      P._parse_pages("99-100", 3) == [1, 2, 3], str(P._parse_pages("99-100", 3)))
check("乱码 → 退回全部", P._parse_pages("abc", 3) == [1, 2, 3])

print("\n== B. 入口/描述层（模型看得到的东西）==")
check("★action 枚举里有 render", "render" in P.TOOL_PARAMETERS["properties"]["action"]["enum"])
check("★schema 补了 pages / dpi",
      "pages" in P.TOOL_PARAMETERS["properties"] and "dpi" in P.TOOL_PARAMETERS["properties"])
check("★TOOL_DESCRIPTION 里告诉了模型 render 怎么用",
      "action=render" in P.TOOL_DESCRIPTION and "PNG" in P.TOOL_DESCRIPTION)
r = json.loads(P.execute(action="render", path=str(ROOT / "_not_exist_193.pptx")))
check("文件不存在 → 如实失败（不抛异常）", r.get("success") is False and "找不到" in r.get("text", ""),
      r.get("text", "")[:60])
r = json.loads(P.execute(action="乱写一个"))
check("不认识的 action 提示里带 render", "render" in r.get("text", ""), r.get("text", "")[:70])
# ★ 1.9.3 第 21 轮真机教训：render 出了 6 张图，模型却在思考里说「我看不到图」，
#   于是 render 白跑、排版根本没核对。所以**必须明确指路**：去看这些 PNG 用 file_ops view。
check("★TOOL_DESCRIPTION 指路「用 file_ops view 看图」（否则模型以为自己瞎）",
      "file_ops view" in P.TOOL_DESCRIPTION, P.TOOL_DESCRIPTION[-120:].replace("\n", " "))

print("\n== D. ★render 一上手就抓到的 create 缺陷：首页被画两遍 ==")
# 旧行为：`# 标题` + `- 要点` 的大纲 → 「封面(只有标题) + 内容页(同一标题+要点)」，
# 标题在片子里出现两次，封面还带页码 0。这里用 python-pptx 直接数页数与标题出现次数。
_out = ROOT / "data" / "桌宠私有data" / "pptx" / "_render_test_193"
_r = json.loads(P.execute(action="create",
                          outline="# 甲标题\n- 甲要点一\n- 甲要点二\n# 乙标题\n- 乙要点",
                          out_dir=str(_out), filename="dup_probe.pptx"))
check("create 成功", _r.get("success") is True, _r.get("text", "")[:80])
check("★页数正确（2 页，不再多出一页）", _r.get("slides") == 2, f"slides={_r.get('slides')}")
from pptx import Presentation as _Prs  # noqa: E402
_p = _Prs(str(_out / "dup_probe.pptx"))
_texts = ["\n".join(sh.text_frame.text for sh in sl.shapes if sh.has_text_frame)
          for sl in _p.slides]
check("★标题只出现一次（不再重复画）",
      sum(t.count("甲标题") for t in _texts) == 1,
      f"出现 {sum(t.count('甲标题') for t in _texts)} 次")
check("★第一页的要点没丢（被吃进封面，而不是被丢弃）",
      "甲要点一" in _texts[0] and "甲要点二" in _texts[0], _texts[0].replace("\n", " / ")[:80])
check("★封面不打页码「0」", "0" not in _texts[0].split("\n"), _texts[0].replace("\n", " | ")[:80])
check("★第二页正常", "乙标题" in _texts[1] and "乙要点" in _texts[1])
_r2 = json.loads(P.execute(action="create", outline="# 只有一页\n- 唯一要点",
                           out_dir=str(_out), filename="one_probe.pptx"))
check("★单页大纲只出 1 页（不会又变回 2 页）", _r2.get("slides") == 1, f"slides={_r2.get('slides')}")
_r3 = json.loads(P.execute(action="create", outline="# 甲\n- 一\n# 乙\n- 二",
                           title="外部标题", out_dir=str(_out), filename="t_probe.pptx"))
check("★显式 title 与首行不一致时：封面用外部标题，内容页全保留",
      _r3.get("slides") == 3, f"slides={_r3.get('slides')}")

print("\n== D2. ★透明 PNG 不能被压成黑底 / 黑白条纹（真机立绘页就是这么翻车的）==")
# 真机证据：`theresa_prts1.png` 是 1496×1496 的 RGBA 立绘，四角 alpha=0 且
# **透明区的 RGB 是垃圾值**。旧代码 `convert("RGB")` 把这些垃圾画了出来 →
# 用户看到的就是"一堆黑白横条"。这里造一张同样性质的图来钉住。
_tp = _out / "alpha_probe.png"
_im = Image.new("RGBA", (400, 400), (255, 255, 255, 0))      # 全透明，RGB=白（垃圾值）
_im.paste((200, 30, 30, 255), (140, 140, 260, 260))          # 中间一块不透明的红
_im.save(_tp)
_r4 = json.loads(P.execute(action="create", outline="# 透明图测试\n- 看四角颜色\n@" + str(_tp),
                           out_dir=str(_out), filename="alpha_probe.pptx"))
check("透明素材建片成功（没被当坏图跳过）",
      _r4.get("success") is True and not _r4.get("skipped_images"), str(_r4.get("skipped_images")))
_crop = P.DEFAULT_OUT / "_tmp" / "alpha_probe_crop.jpg"
check("裁剪产物落盘", _crop.exists(), str(_crop))
if _crop.exists():
    _im2 = Image.open(_crop).convert("RGB")
    _w, _h = _im2.size
    _corners = [_im2.getpixel(p) for p in ((0, 0), (_w - 1, 0), (0, _h - 1), (_w - 1, _h - 1))]
    _allwhite = all(all(c >= 240 for c in px) for px in _corners)
    check("★四角是**白**底（透明被压到页面底色），不是黑", _allwhite, str(_corners))
    _mid = _im2.getpixel((_w // 2, _h // 2))
    check("★中间那块不透明的内容还在（没被底色盖掉）",
          _mid[0] > 150 and _mid[1] < 90 and _mid[2] < 90, str(_mid))

print("\n== D3. ★1.9.3 第 23 轮：图片不再被硬裁成正方形 + 开头的 @图片 不再开空页 ==")
# 真机证据：模型给的素材 `news_whale.png` 是 690×1215（竖图，比例 0.57），
# 而旧槽位是**写死的 4.6×4.6 正方形** → 上下各切 22%，立绘的头直接被切掉。
_bw, _bh = P._fit_box(690, 1215, 4.45, 5.05)
check("★自适应框保持原图比例（不再拍成正方形）",
      abs((_bw / _bh) - (690 / 1215)) < 0.01, f"{_bw:.2f}x{_bh:.2f} 比例={_bw/_bh:.3f}")
check("★竖图在可用区域里被放大到贴边（高度吃满 5.05）", abs(_bh - 5.05) < 0.02, f"h={_bh:.3f}")
_w2, _h2 = P._fit_box(1600, 800, 4.45, 5.05)
check("★横图也一样：比例保持不变", abs((_w2 / _h2) - 2.0) < 0.01, f"{_w2:.2f}x{_h2:.2f}")

_sl, _lead = P._parse_outline("@cover.png\n\n# 标题甲\n- 要点一\n")
check("★开头的 @图片 不再单独开页（改记为封面图）",
      len(_sl) == 1 and _sl[0]["title"] == "标题甲" and _lead == ["cover.png"],
      f"slides={len(_sl)} leading={_lead}")
_sl2, _lead2 = P._parse_outline("# 甲\n- 一\n@after.png\n")
check("标题之后的 @图片 仍挂当前页", _sl2[0]["imgs"] == ["after.png"] and not _lead2)

_pv = _out / "portrait_probe.png"
Image.new("RGB", (300, 500), (30, 90, 160)).save(_pv)
# ★ 注意：图片挂在**第一页**时会被当封面图（tag=_cover）。这里多给一个 `# 封面`
#   让竖图落到**内容页**上，才测得到内容页的槽位。
_r5 = json.loads(P.execute(action="create",
                           outline="# 封面\n# 竖图页\n- 看它还切不切头\n@" + str(_pv),
                           out_dir=str(_out), filename="portrait_probe.pptx"))
check("竖图页拍完（共 2 页）", _r5.get("slides") == 2, str(_r5.get("slides")))
_crop2 = P.DEFAULT_OUT / "_tmp" / "portrait_probe_p0.jpg"
check("竖图裁剪产物落盘", _crop2.exists(), str(_crop2))
if _crop2.exists():
    with Image.open(_crop2) as _im4:
        check("★★裁剪产物仍是竖构图（旧代码这里会变成 1:1 正方形）",
              _im4.height > _im4.width and abs((_im4.width / _im4.height) - 0.6) < 0.02,
              f"{_im4.size} 比例={_im4.width/_im4.height:.3f}")

print("\n== D4. ★字号自适应（防文本框溢出）+ 主题底色 ==")
# 6 条、每条 96 个字：旧代码只看"条数≤6 → 18pt"，必然溢出文本框；
# 新代码按估算行数缩号。
_long = ["这是一条特别特别长的要点" * 8] * 6
_old_would_be = 18 if len(_long) <= 6 else 14
_fitted = P._fit_body_size(_long, 7.1, 5.0, 18)
check("★超长要点会自动缩号（旧代码只看条数 → 还是 18pt，会撑爆文本框）",
      _fitted < _old_would_be, f"新={_fitted}pt 旧={_old_would_be}pt")
check("短要点不缩号（保持 18pt）", P._fit_body_size(["短", "也很短"], 11.9, 5.0, 18) == 18)

_rd = json.loads(P.execute(action="create", outline="# 深色主题\n- 底色应该被画出来",
                           out_dir=str(_out), filename="dark_probe.pptx", theme="dark"))
check("dark 主题可用", _rd.get("success") is True, str(_rd.get("text"))[:60])
_prs2 = _Prs(str(_out / "dark_probe.pptx"))
_bg_ok = False
for _sh in _prs2.slides[0].shapes:
    try:
        if _sh.fill.type is not None and _sh.fill.fore_color.rgb == P._rgb((0x12, 0x10, 0x16)):
            _bg_ok = True
    except Exception:
        pass
check("★dark 主题真的画了底色（以前主题只改字色，永远是一张白纸）", _bg_ok)

_sizes = []
for _sh in _prs2.slides[0].shapes:
    if _sh.has_text_frame:
        for _p in _sh.text_frame.paragraphs:
            for _r in _p.runs:
                if _r.font.size:
                    _sizes.append(_r.font.size.pt)
check("★封面用大标题（44pt），与内容页区分开",
      max(_sizes or [0]) >= 40, str(sorted(set(_sizes))))

print("\n== E. ★1.9.3 第 24 轮：素材体检（横幅卡要提醒）+ 定位改回「以自己写脚本为主」==")
# 真机证据：模型给铃兰那套片子配图时用了 `suzuran_3.jpg`(1013×427) 与
# `suzuran_b.png`(800×420) —— 都是官方**干员档案横幅卡**，本来就只画到腰/胸。
# 工具原样插进去（尺寸与源文件逐个比对过，没裁），但用户看到的就是"人被切了"。
check("横幅卡要被点名（1013×427 → 比例 2.37）",
      bool(P._asset_check("suzuran_3.jpg", 1013, 427))
      and "横幅卡" in P._asset_check("suzuran_3.jpg", 1013, 427),
      str(P._asset_check("suzuran_3.jpg", 1013, 427)))
check("800×420（比例 1.905）也要点名",
      bool(P._asset_check("suzuran_b.png", 800, 420)),
      str(P._asset_check("suzuran_b.png", 800, 420)))
check("正常竖图不吵（690×1215 → 比例 0.57）",
      P._asset_check("news_whale.png", 690, 1215) is None)
check("正常方图不吵（700×700）", P._asset_check("lisa.png", 700, 700) is None)
check("超长竖条要提醒（400×1200 → 0.33）",
      bool(P._asset_check("tall.png", 400, 1200)),
      str(P._asset_check("tall.png", 400, 1200)))
check("小图要提醒（300×300）", P._asset_check("small.png", 300, 300) is not None)

_banner = _out / "banner_probe.jpg"
Image.new("RGB", (1013, 427), (40, 40, 60)).save(_banner)
_rw = json.loads(P.execute(action="create",
                           outline="# 横幅卡测试\n- 这张是横幅卡\n@" + str(_banner),
                           out_dir=str(_out), filename="banner_probe.pptx"))
check("★create 返回里带 asset_warnings（不只放在正文里）",
      isinstance(_rw.get("asset_warnings"), list) and len(_rw["asset_warnings"]) >= 1,
      str(_rw.get("asset_warnings"))[:120])
check("★返回文案里也写了体检结论（模型读的是 text）",
      "素材体检" in _rw.get("text", "") and "横幅卡" in _rw.get("text", ""),
      _rw.get("text", "")[-120:].replace("\n", " "))

_lc = (ROOT / "src" / "core" / "llm_client.py").read_text(encoding="utf-8")
check("★系统提示词不再强制「必须用 pptx 工具、不要自己写脚本」",
      "不要**自己写 python-pptx 脚本" not in _lc and "**不要**自己写 python-pptx 脚本" not in _lc)
check("★系统提示词改回「想要设计感就自己写脚本」",
      "自己写 python-pptx 脚本" in _lc and "有设计感" in _lc)
check("★系统提示词要求做完用 render + file_ops view 自检",
      "action=render" in _lc and "file_ops view" in _lc)
check("★系统提示词提醒分辨横幅卡（宽高比 ≥1.9）",
      "1.9" in _lc and "横幅卡" in _lc)
check("★工具描述也不再自称「优先用它、不要自己写脚本」",
      "**做 PPT 优先用它，不要自己写脚本**" not in P.TOOL_DESCRIPTION
      and "自己写 python-pptx 脚本更好" in P.TOOL_DESCRIPTION)


def _powerpnt_alive():
    try:
        import psutil
        return any((p.info.get("name") or "").lower() == "powerpnt.exe"
                   for p in psutil.process_iter(["name"]))
    except Exception:
        return False


def _wait_no_powerpnt(timeout=20.0):
    """等 POWERPNT 真的退出（COM 的 Quit 是异步的，进程会滞留一两秒）。"""
    import time
    t0 = time.time()
    while time.time() - t0 < timeout:
        if not _powerpnt_alive():
            return True
        time.sleep(0.5)
    return False


# 先判定"有没有 Office"，用一次最轻的 Dispatch；失败就 SKIP（别的机器上不该红）
_has_office = False
try:
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()
    try:
        _app = win32com.client.Dispatch("PowerPoint.Application")
        _has_office = True
        try:
            _app.Quit()
        except Exception:
            pass
    finally:
        pythoncom.CoUninitialize()
except Exception as e:
    print(f"  [SKIP] 本机没有可用的 PowerPoint COM：{e}")

if not _has_office:
    print("\n★ 本机没有 PowerPoint → 真机层跳过（不算失败；render 会如实返回失败）")
else:
    _wait_no_powerpnt()          # 把探测那次的进程等干净，再取基线
    print("\n== C. 真机：真的导出一份 3 页 PPT 来看看 ==")
    from pptx import Presentation
    from pptx.util import Inches

    tmp_dir = ROOT / "data" / "桌宠私有data" / "pptx" / "_render_test_193"
    tmp_dir.mkdir(parents=True, exist_ok=True)
    src = tmp_dir / "probe.pptx"
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for i in range(1, 4):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        tb = s.shapes.add_textbox(Inches(0.6), Inches(0.5), Inches(12), Inches(1))
        tb.text_frame.text = f"第 {i} 页探针"
    prs.save(str(src))
    check("探针 pptx 已生成（3 页）", src.exists())

    before = _powerpnt_alive()
    out_dir = tmp_dir / "png"
    r = json.loads(P.execute(action="render", path=str(src), out_dir=str(out_dir), dpi=96))
    check("★render 成功返回", r.get("success") is True, r.get("text", "")[:100])
    imgs = [Path(x) for x in r.get("images", [])]
    check("★3 页全导出，且文件真的落盘",
          len(imgs) == 3 and all(p.exists() and p.stat().st_size > 0 for p in imgs),
          f"{len(imgs)} 张")
    try:
        from PIL import Image
        ok_decode = True
        for p in imgs:
            with Image.open(p) as im:
                im.verify()
        check("★导出的 PNG 都能解码（不是 0 字节残片）", ok_decode)
    except Exception as e:
        check("★导出的 PNG 都能解码（不是 0 字节残片）", False, str(e))

    r2 = json.loads(P.execute(action="render", path=str(src), out_dir=str(out_dir),
                              pages="2", dpi=96))
    imgs2 = [Path(x) for x in r2.get("images", [])]
    check("★pages=\"2\" 只导第 2 页", len(imgs2) == 1 and "_p02" in imgs2[0].name,
          str([p.name for p in imgs2]))
    check("★text 里给了图片绝对路径（模型要照着去看）",
          all(p.name in r2.get("text", "") for p in imgs2))
    check("★render 成功返回的文案里就指路 file_ops view（模型读的是返回值）",
          "file_ops view" in r2.get("text", ""), r2.get("text", "")[-90:].replace("\n", " "))

    # ★ 真机踩到的坑：PowerPoint COM **不吃相对路径**（报 -2147024893 找不到路径）。
    #   回归里用相对路径再导一次，防止以后有人把 resolve() 去掉。
    os.chdir(ROOT)
    rel = os.path.relpath(str(src), str(ROOT))
    r3 = json.loads(P.execute(action="render", path=rel, out_dir=str(out_dir), pages="1", dpi=96))
    check("★相对路径也能导（COM 需要绝对路径，工具内部必须转）",
          r3.get("success") is True, r3.get("text", "")[:90])
    check("★返回的图片路径是绝对的（模型/后续工具不必猜 cwd）",
          all(os.path.isabs(x) for x in r3.get("images", [])) and bool(r3.get("images")),
          str(r3.get("images", [])[:1]))

    check("★跑完不留下孤儿 POWERPNT（我们开的我们关；Quit 是异步的所以等一下）",
          _wait_no_powerpnt() if not before else _powerpnt_alive(),
          f"before={before}")


print(f"\n==== 结果: {len(PASS)} PASS / {len(FAIL)} FAIL ====")
for f in FAIL:
    print("  FAIL:", f)

# 全绿就把探针目录清掉（失败时**留着**，方便去看导出的 PNG 差在哪）
if not FAIL:
    try:
        import shutil
        shutil.rmtree(ROOT / "data" / "桌宠私有data" / "pptx" / "_render_test_193",
                      ignore_errors=True)
    except Exception:
        pass

sys.exit(1 if FAIL else 0)
